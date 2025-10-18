#!/usr/bin/env python3
"""
Diagnostic script to test B01 connection with different settings.
"""

from netmiko import ConnectHandler
from openpyxl import load_workbook
import getpass
import os
import sys

# Configuration
script_dir = os.path.dirname(os.path.abspath(__file__))
input_file = os.path.join(script_dir, "network_inputs.xlsx")

# Default credentials
default_username = "NCMSOLK"
default_password = "mhb5N2Ap"

def test_connection(device_ip, device_type, username, password, timeout=15):
    """Test connection with detailed error reporting."""
    print(f"\n{'='*70}")
    print(f"Testing: {device_type} on {device_ip}")
    print(f"Username: {username}")
    print(f"Timeout: {timeout}s")
    print(f"{'='*70}")
    
    try:
        print("🔗 Attempting connection...")
        conn = ConnectHandler(
            device_type=device_type,
            ip=device_ip,
            username=username,
            password=password,
            timeout=timeout,
            verbose=True  # Enable verbose output
        )
        print("✅ CONNECTION SUCCESSFUL!")
        
        # Try a simple command
        print("\n▶ Testing command execution: show version")
        output = conn.send_command("show version", read_timeout=10)
        print(f"✅ Command executed successfully ({len(output)} chars)")
        print(f"First 200 chars:\n{output[:200]}")
        
        conn.disconnect()
        print("\n✅ Disconnected successfully")
        return True
        
    except Exception as e:
        print(f"\n❌ CONNECTION FAILED!")
        print(f"Error Type: {type(e).__name__}")
        print(f"Error Message: {str(e)}")
        print(f"\nFull Error Details:")
        print(f"{repr(e)}")
        return False

def main():
    # Load Excel to find B01
    if not os.path.exists(input_file):
        print(f"❌ Excel file not found: {input_file}")
        print("Creating test without Excel - you'll need to enter B01 IP manually")
        b01_ip = input("Enter B01 IP address: ")
    else:
        wb = load_workbook(input_file)
        devices_ws = wb["Devices"]
        
        b01_ip = None
        for row in devices_ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and "B01" in row[0]:
                b01_ip = row[1]
                print(f"✅ Found B01 in Excel: {row[0]} ({b01_ip})")
                break
        
        if not b01_ip:
            print("❌ No B01 device found in Excel")
            b01_ip = input("Enter B01 IP address manually: ")
    
    print("\n" + "="*70)
    print("DIAGNOSTIC TEST FOR B01 CONNECTION")
    print("="*70)
    print(f"Target Device: {b01_ip}")
    print(f"Default Username: {default_username}")
    print(f"Default Password: {'*' * len(default_password)}")
    
    # Test with different device types
    device_types_to_test = [
        "cisco_nxos",
        "cisco_ios",
        "cisco_xe",
        "cisco_xr",
    ]
    
    print("\n🔬 Testing with DEFAULT credentials...")
    username = default_username
    password = default_password
    
    for device_type in device_types_to_test:
        result = test_connection(b01_ip, device_type, username, password)
        if result:
            print(f"\n🎉 SUCCESS with device_type: {device_type}")
            break
    
    # If all failed, try with manual credentials
    if not any([test_connection(b01_ip, dt, username, password, timeout=5) 
                for dt in device_types_to_test]):
        print("\n" + "="*70)
        print("❌ All device types failed with default credentials")
        print("="*70)
        
        print("\nWould you like to try with different credentials? (y/n)")
        choice = input("> ").lower()
        
        if choice == 'y':
            username = input("Enter username: ")
            password = getpass.getpass("Enter password: ")
            
            print("\n🔬 Testing with PROVIDED credentials...")
            for device_type in device_types_to_test:
                result = test_connection(b01_ip, device_type, username, password)
                if result:
                    print(f"\n🎉 SUCCESS with device_type: {device_type}")
                    break
    
    print("\n" + "="*70)
    print("DIAGNOSTIC TEST COMPLETE")
    print("="*70)
    print("\nRecommendations based on results above:")
    print("1. Check the error messages to identify the issue")
    print("2. Verify the device_type that works (if any)")
    print("3. Update your Excel file or script accordingly")
    print("4. Check network connectivity if you see timeouts")
    print("5. Verify SSH is enabled on the device")

if __name__ == "__main__":
    main()
