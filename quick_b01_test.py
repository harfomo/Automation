#!/usr/bin/env python3
"""
Quick B01 connection test - finds the right device type.
"""

from netmiko import ConnectHandler
from openpyxl import load_workbook
import os

# Configuration
input_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "network_inputs.xlsx")

# Credentials
username = "NCMSOLK"
password = "mhb5N2Ap"

# Load B01 from Excel
print("📂 Loading B01 from Excel...")
wb = load_workbook(input_file)
devices_ws = wb["Devices"]

b01_ip = None
for row in devices_ws.iter_rows(min_row=2, values_only=True):
    if row and row[0] and "B01" in row[0]:
        b01_ip = row[1]
        print(f"✅ Found B01: {row[0]} at {b01_ip}\n")
        break

if not b01_ip:
    print("❌ B01 not found in Excel")
    exit(1)

# Test different device types
device_types = [
    "cisco_nxos",
    "cisco_ios", 
    "cisco_xe",
    "cisco_xr",
    "cisco_asa",
]

print("="*70)
print("TESTING DIFFERENT DEVICE TYPES")
print("="*70)

for dtype in device_types:
    print(f"\n{'─'*70}")
    print(f"Testing: {dtype}")
    print(f"{'─'*70}")
    
    try:
        print(f"🔗 Connecting...")
        conn = ConnectHandler(
            device_type=dtype,
            ip=b01_ip,
            username=username,
            password=password,
            timeout=20,
            session_log='b01_session.log'  # Log session for debugging
        )
        
        print(f"✅ CONNECTION SUCCESSFUL with {dtype}!")
        
        # Try to get prompt
        prompt = conn.find_prompt()
        print(f"✅ Device prompt: {prompt}")
        
        # Try a simple command
        print(f"▶ Testing command: show version")
        output = conn.send_command("show version", read_timeout=10)
        print(f"✅ Command successful! Output length: {len(output)} chars")
        print(f"First 150 chars of output:")
        print(f"  {output[:150]}")
        
        # Try the BGP command
        print(f"\n▶ Testing BGP command: sh run bgp | i \"router bgp\"")
        bgp_output = conn.send_command("sh run bgp | i \"router bgp\"", read_timeout=10)
        print(f"✅ BGP command successful!")
        print(f"Output: {bgp_output}")
        
        conn.disconnect()
        
        print(f"\n{'🎉'*35}")
        print(f"SUCCESS! Use device_type: {dtype}")
        print(f"{'🎉'*35}")
        print(f"\nUpdate your Excel file:")
        print(f"  | Cilli_Hostname | IP | Device_Type |")
        print(f"  | NWCSDEBGB01    | {b01_ip} | {dtype} |")
        break
        
    except Exception as e:
        print(f"❌ Failed with {dtype}")
        print(f"   Error: {type(e).__name__}")
        print(f"   Message: {str(e)[:200]}")
        continue

print(f"\n{'='*70}")
print("DIAGNOSTIC COMPLETE")
print("="*70)
print("\nCheck b01_session.log for full session details if needed")
