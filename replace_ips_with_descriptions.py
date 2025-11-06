from openpyxl import load_workbook
import os
import re

# =========================================================
# Configuration
# =========================================================
script_dir = os.path.dirname(os.path.abspath(__file__))
network_inputs_file = os.path.join(script_dir, "network_inputs.xlsx")
mop_temp_file = os.path.join(script_dir, "MOPtemp.xlsx")
output_file = os.path.join(script_dir, "MOPtemp_updated.xlsx")

# =========================================================
# Validation
# =========================================================
if not os.path.exists(network_inputs_file):
    print(f"❌ Error: '{network_inputs_file}' not found!")
    print("   Please run the network_device_collector.py script first.")
    raise SystemExit

if not os.path.exists(mop_temp_file):
    print(f"❌ Error: '{mop_temp_file}' not found!")
    print("   Please ensure MOPtemp.xlsx exists in the script directory.")
    raise SystemExit

# =========================================================
# Load network_inputs.xlsx and build IP -> Description mapping
# =========================================================
print("📖 Loading network_inputs.xlsx...")
network_wb = load_workbook(network_inputs_file, data_only=True)

if "Neighbors" not in network_wb.sheetnames:
    print("❌ Error: 'Neighbors' sheet not found in network_inputs.xlsx")
    raise SystemExit

neighbors_sheet = network_wb["Neighbors"]

# Build mapping: Neighbor IP -> Neighbor Description
ip_to_description = {}
replacement_count = 0

print("🔍 Building IP to Description mapping from Neighbors sheet...")
for row in neighbors_sheet.iter_rows(min_row=2, values_only=True):
    if not row or len(row) < 3:
        continue
    
    # Column structure: [Device IP, Neighbor Description, Neighbor IP, ...]
    neighbor_desc = row[1]  # Column B: Neighbor Description
    neighbor_ip = row[2]    # Column C: Neighbor IP
    
    if neighbor_ip and neighbor_desc:
        neighbor_ip_str = str(neighbor_ip).strip()
        neighbor_desc_str = str(neighbor_desc).strip()
        
        if neighbor_ip_str and neighbor_desc_str:
            ip_to_description[neighbor_ip_str] = neighbor_desc_str

print(f"✅ Found {len(ip_to_description)} unique Neighbor IP -> Description mappings")

if not ip_to_description:
    print("⚠️  No valid mappings found. Exiting.")
    raise SystemExit

# =========================================================
# Load MOPtemp.xlsx and search/replace
# =========================================================
print(f"\n📖 Loading {mop_temp_file}...")
mop_wb = load_workbook(mop_temp_file)

print(f"📝 Searching and replacing IPs in all sheets...")

for sheet_name in mop_wb.sheetnames:
    sheet = mop_wb[sheet_name]
    print(f"\n  🔎 Processing sheet: '{sheet_name}'")
    sheet_replacements = 0
    
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            
            cell_value = str(cell.value)
            original_value = cell_value
            
            # Check if any of our IPs exist in this cell
            for ip, description in ip_to_description.items():
                # Use word boundary matching to avoid partial replacements
                # For IPv6 addresses with colons, we need to be careful with regex
                # Try exact match first, then pattern match
                
                if ip in cell_value:
                    # Replace the IP with the description
                    # Use whole word replacement to avoid replacing part of text
                    pattern = re.escape(ip)
                    cell_value = re.sub(r'\b' + pattern + r'\b', description, cell_value)
            
            # If the value changed, update the cell
            if cell_value != original_value:
                cell.value = cell_value
                sheet_replacements += 1
                replacement_count += 1
                print(f"    ✓ Replaced in cell {cell.coordinate}: '{original_value}' -> '{cell_value}'")
    
    if sheet_replacements > 0:
        print(f"  ✅ Made {sheet_replacements} replacement(s) in '{sheet_name}'")
    else:
        print(f"  ℹ️  No replacements needed in '{sheet_name}'")

# =========================================================
# Save the updated workbook
# =========================================================
print(f"\n💾 Saving updated workbook to '{output_file}'...")
mop_wb.save(output_file)

print(f"\n✅ Script completed successfully!")
print(f"📊 Total replacements made: {replacement_count}")
print(f"📁 Updated file saved as: {output_file}")
print(f"\n💡 Tip: Review the changes and if satisfied, you can rename '{output_file}' to 'MOPtemp.xlsx'")
