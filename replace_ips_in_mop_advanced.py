from openpyxl import load_workbook
import os
import re
import argparse

# =========================================================
# Configuration & Arguments
# =========================================================
script_dir = os.path.dirname(os.path.abspath(__file__))
network_inputs_file = os.path.join(script_dir, "network_inputs.xlsx")
mop_temp_file = os.path.join(script_dir, "MOPtemp.xlsx")

parser = argparse.ArgumentParser(
    description='Replace Neighbor IPs with Descriptions in MOPtemp.xlsx using mappings from network_inputs.xlsx'
)
parser.add_argument('--overwrite', action='store_true', 
                    help='Overwrite the original MOPtemp.xlsx instead of creating a new file')
parser.add_argument('--sheets', nargs='+', 
                    help='Specific sheet names to process in MOPtemp.xlsx (default: all sheets)')
parser.add_argument('--preview', action='store_true',
                    help='Preview changes without saving (dry run)')
parser.add_argument('--verbose', action='store_true',
                    help='Show detailed output for every replacement')
parser.add_argument('--network-file', default=None,
                    help='Path to network inputs file (default: network_inputs.xlsx in script directory)')
parser.add_argument('--mop-file', default=None,
                    help='Path to MOP template file (default: MOPtemp.xlsx in script directory)')
args = parser.parse_args()

# Override file paths if provided
if args.network_file:
    network_inputs_file = args.network_file
if args.mop_file:
    mop_temp_file = args.mop_file

# =========================================================
# Validation
# =========================================================
if not os.path.exists(network_inputs_file):
    print(f"❌ Error: '{network_inputs_file}' not found!")
    print("   Please run the network_device_collector.py script first.")
    raise SystemExit

if not os.path.exists(mop_temp_file):
    print(f"❌ Error: '{mop_temp_file}' not found!")
    print("   Please ensure MOPtemp.xlsx exists in the script directory.")
    raise SystemExit

# =========================================================
# Load network_inputs.xlsx and build IP -> Description mapping
# =========================================================
print("📖 Loading network_inputs.xlsx...")
print(f"   Source: {network_inputs_file}")
network_wb = load_workbook(network_inputs_file, data_only=True)

if "Neighbors" not in network_wb.sheetnames:
    print("❌ Error: 'Neighbors' sheet not found in network_inputs.xlsx")
    print(f"   Available sheets: {', '.join(network_wb.sheetnames)}")
    raise SystemExit

neighbors_sheet = network_wb["Neighbors"]

# Build mapping: Neighbor IP -> Neighbor Description
ip_to_description = {}
replacement_count = 0
changes_log = []

print("🔍 Building IP to Description mapping from Neighbors sheet...")

# Try to find the header row to determine column positions
header_row = None
neighbor_desc_col = None
neighbor_ip_col = None
device_ip_col = None

for row_idx, row in enumerate(neighbors_sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
    if row:
        for col_idx, cell_value in enumerate(row):
            if cell_value:
                cell_str = str(cell_value).strip().lower()
                if 'device ip' in cell_str:
                    device_ip_col = col_idx
                    header_row = row_idx
                elif 'neighbor description' in cell_str or 'description' in cell_str:
                    neighbor_desc_col = col_idx
                    header_row = row_idx
                elif 'neighbor ip' in cell_str or (cell_str == 'neighbor ip'):
                    neighbor_ip_col = col_idx
                    header_row = row_idx
        
        if neighbor_desc_col is not None and neighbor_ip_col is not None:
            break

# Default to standard layout if headers not found
if neighbor_desc_col is None or neighbor_ip_col is None:
    print("   ⚠️  Could not auto-detect columns, using default layout:")
    print("      Column A (index 0) = Device IP")
    print("      Column B (index 1) = Neighbor Description")
    print("      Column C (index 2) = Neighbor IP")
    device_ip_col = 0
    neighbor_desc_col = 1
    neighbor_ip_col = 2
    header_row = 1
else:
    print(f"   ✅ Auto-detected columns:")
    if device_ip_col is not None:
        print(f"      Column {chr(65 + device_ip_col)} = Device IP")
    print(f"      Column {chr(65 + neighbor_desc_col)} = Neighbor Description")
    print(f"      Column {chr(65 + neighbor_ip_col)} = Neighbor IP")

# Build the mapping
duplicate_count = 0
for row in neighbors_sheet.iter_rows(min_row=header_row + 1, values_only=True):
    if not row or len(row) <= max(neighbor_desc_col, neighbor_ip_col):
        continue
    
    neighbor_desc = row[neighbor_desc_col] if neighbor_desc_col < len(row) else None
    neighbor_ip = row[neighbor_ip_col] if neighbor_ip_col < len(row) else None
    
    if neighbor_ip and neighbor_desc:
        neighbor_ip_str = str(neighbor_ip).strip()
        neighbor_desc_str = str(neighbor_desc).strip()
        
        # Skip empty or placeholder values
        if neighbor_ip_str and neighbor_desc_str and neighbor_ip_str.lower() not in ['none', 'n/a', '-', 'null', 'error']:
            # Store mapping (if duplicate, keep first occurrence)
            if neighbor_ip_str in ip_to_description:
                duplicate_count += 1
                if args.verbose:
                    print(f"   ⚠️  Duplicate IP: {neighbor_ip_str}")
                    print(f"      Existing: {ip_to_description[neighbor_ip_str]}")
                    print(f"      Skipping: {neighbor_desc_str}")
            else:
                ip_to_description[neighbor_ip_str] = neighbor_desc_str

print(f"✅ Found {len(ip_to_description)} unique Neighbor IP -> Description mappings")
if duplicate_count > 0:
    print(f"   ℹ️  Skipped {duplicate_count} duplicate IPs (kept first occurrence)")

if not ip_to_description:
    print("⚠️  No valid mappings found. Exiting.")
    raise SystemExit

# Show mappings preview
print("\n📋 IP to Description mappings:")
if len(ip_to_description) <= 20 or args.verbose:
    for ip, desc in ip_to_description.items():
        print(f"   {ip} → {desc}")
else:
    for i, (ip, desc) in enumerate(list(ip_to_description.items())[:10]):
        print(f"   {ip} → {desc}")
    print(f"   ... and {len(ip_to_description) - 10} more")

# Close network_inputs workbook
network_wb.close()

# =========================================================
# Helper functions for replacements
# =========================================================
def replace_ip_in_text(text, ip, description):
    """
    Replace IP with description in text, handling both IPv4 and IPv6.
    Returns (new_text, was_replaced)
    """
    original_text = text
    
    # Create case-insensitive pattern
    pattern = re.compile(re.escape(ip), re.IGNORECASE)
    
    # Check if pattern exists
    if not pattern.search(text):
        return text, False
    
    # For IPv6 (contains colons) and IPv4, use different strategies
    if ':' in ip:
        # IPv6 - direct replacement
        text = pattern.sub(description, text)
    else:
        # IPv4 - try word boundary first
        try:
            pattern_wb = re.compile(r'\b' + re.escape(ip) + r'\b', re.IGNORECASE)
            if pattern_wb.search(text):
                text = pattern_wb.sub(description, text)
            else:
                text = pattern.sub(description, text)
        except:
            text = pattern.sub(description, text)
    
    return text, (text != original_text)

def replace_route_counts_in_next_column(next_cell_value, description, route_type, is_evpn=False):
    """
    Replace route counts in the next column with description-based labels.
    
    For regular BGP routes:
      - advertised-routes: <number> routes -> {description}_adv_routes routes
      - received-routes: <number> routes -> {description}_recv_routes routes
    
    For EVPN routes:
      - advertised-routes evpn: (entire cell content) -> {description}_BGP_EVPN_Prefix_Adv
      - received-routes evpn: (entire cell content) -> {description}_BGP_EVPN_Prefix_Recv
    
    Args:
        next_cell_value: Value of the next column cell
        description: The neighbor description to use
        route_type: 'advertised' or 'received'
        is_evpn: True if this is an EVPN route command
    
    Returns (new_text, was_replaced)
    """
    if not next_cell_value or not description:
        return next_cell_value, False
    
    text = str(next_cell_value).strip()
    original_text = text
    
    # Handle EVPN routes - replace entire cell content with simple label
    if is_evpn:
        # EVPN routes have complex format like:
        # "Auto-Disc-0 routes, IP-Prefix-2471 routes, IPv6-Prefix-8884 routes"
        # Replace entire content with description-based label
        if route_type == 'advertised':
            new_text = f"{description}_BGP_EVPN_Prefix_Adv"
        elif route_type == 'received':
            new_text = f"{description}_BGP_EVPN_Prefix_Recv"
        else:
            return text, False
        
        # Only replace if there's actually content to replace
        if text and text.lower() not in ['none', 'n/a', '-', 'null']:
            return new_text, True
        return text, False
    
    # Handle regular BGP routes - match specific pattern
    # Pattern to match: optional whitespace, digits, optional whitespace, "routes"
    # Examples: "311 routes", "  500 routes  ", "123routes"
    pattern = re.compile(r'^\s*(\d+)\s*(routes?)\s*$', re.IGNORECASE)
    match = pattern.match(text)
    
    if match:
        if route_type == 'advertised':
            new_text = f"{description}_adv_routes routes"
        elif route_type == 'received':
            new_text = f"{description}_recv_routes routes"
        else:
            return text, False
        
        return new_text, True
    
    return text, False

# =========================================================
# Load MOPtemp.xlsx and search/replace
# =========================================================
print(f"\n📖 Loading {mop_temp_file}...")
print(f"   Target: {mop_temp_file}")
mop_wb = load_workbook(mop_temp_file)

# Determine which sheets to process
if args.sheets:
    sheets_to_process = args.sheets
    # Validate sheet names
    invalid_sheets = [s for s in sheets_to_process if s not in mop_wb.sheetnames]
    if invalid_sheets:
        print(f"⚠️  Warning: The following sheets were not found: {', '.join(invalid_sheets)}")
        sheets_to_process = [s for s in sheets_to_process if s in mop_wb.sheetnames]
else:
    sheets_to_process = mop_wb.sheetnames

if not sheets_to_process:
    print("❌ No valid sheets to process. Exiting.")
    raise SystemExit

print(f"📝 Searching and replacing IPs in {len(sheets_to_process)} sheet(s) of MOPtemp.xlsx...")
if args.preview:
    print("🔍 PREVIEW MODE - No changes will be saved\n")

for sheet_name in sheets_to_process:
    sheet = mop_wb[sheet_name]
    print(f"\n  🔎 Processing sheet: '{sheet_name}'")
    sheet_replacements = 0
    
    for row in sheet.iter_rows():
        # Convert row to list so we can access cells by index
        row_cells = list(row)
        
        for cell_idx, cell in enumerate(row_cells):
            if cell.value is None:
                continue
            
            cell_value = str(cell.value)
            original_value = cell_value
            replaced_ips = []
            descriptions_used = []
            
            # Step 1: Replace IPs with descriptions
            for ip, description in ip_to_description.items():
                new_value, was_replaced = replace_ip_in_text(cell_value, ip, description)
                if was_replaced:
                    cell_value = new_value
                    replaced_ips.append((ip, description))
                    descriptions_used.append(description)
            
            # Step 2: Check if this cell has advertised-routes or received-routes
            # If yes, replace route count in the NEXT column
            route_count_replaced = False
            next_cell_modified = False
            next_cell_coord = None
            if descriptions_used:
                description = descriptions_used[0]
                route_type = None
                is_evpn = False
                
                # Check for EVPN routes first (more specific)
                if 'advertised-routes' in cell_value.lower() and 'evpn' in cell_value.lower():
                    route_type = 'advertised'
                    is_evpn = True
                elif 'received-routes' in cell_value.lower() and 'evpn' in cell_value.lower():
                    route_type = 'received'
                    is_evpn = True
                # Then check for regular BGP routes
                elif 'advertised-routes' in cell_value.lower():
                    route_type = 'advertised'
                elif 'received-routes' in cell_value.lower():
                    route_type = 'received'
                
                # If we found a route command, check the next column
                if route_type and cell_idx + 1 < len(row_cells):
                    next_cell = row_cells[cell_idx + 1]
                    if next_cell.value is not None:
                        new_next_value, next_cell_modified = replace_route_counts_in_next_column(
                            next_cell.value, description, route_type, is_evpn
                        )
                        if next_cell_modified:
                            if not args.preview:
                                next_cell.value = new_next_value
                            route_count_replaced = True
                            next_cell_coord = next_cell.coordinate
            
            # If the value changed, update the cell
            if cell_value != original_value:
                if not args.preview:
                    cell.value = cell_value
                
                sheet_replacements += 1
                replacement_count += 1
                
                change_info = {
                    'sheet': sheet_name,
                    'cell': cell.coordinate,
                    'original': original_value,
                    'new': cell_value,
                    'replacements': replaced_ips,
                    'route_count_replaced': route_count_replaced,
                    'next_cell_coord': next_cell_coord
                }
                changes_log.append(change_info)
                
                # Show details based on verbosity
                if args.verbose or sheet_replacements <= 5:
                    print(f"    ✓ Cell {cell.coordinate}:")
                    print(f"      Before: {original_value[:70]}{'...' if len(original_value) > 70 else ''}")
                    print(f"      After:  {cell_value[:70]}{'...' if len(cell_value) > 70 else ''}")
                    for ip, desc in replaced_ips:
                        print(f"      • Replaced IP: {ip} → {desc}")
                    if route_count_replaced and next_cell_coord:
                        print(f"      • Replaced route count in next column ({next_cell_coord})")
    
    if sheet_replacements > 0:
        status = "Would make" if args.preview else "Made"
        print(f"  ✅ {status} {sheet_replacements} replacement(s) in '{sheet_name}'")
        if not args.verbose and sheet_replacements > 5:
            print(f"     (showing first 5, use --verbose to see all)")
    else:
        print(f"  ℹ️  No replacements needed in '{sheet_name}'")

# =========================================================
# Save the updated workbook
# =========================================================
if not args.preview:
    if args.overwrite:
        output_file = mop_temp_file
        print(f"\n💾 Saving changes to original file '{output_file}'...")
    else:
        output_file = os.path.join(script_dir, "MOPtemp_updated.xlsx")
        print(f"\n💾 Saving updated workbook to '{output_file}'...")
    
    mop_wb.save(output_file)
    print(f"✅ File saved successfully!")
else:
    print(f"\n🔍 Preview complete - no changes were saved")

# =========================================================
# Summary
# =========================================================
print(f"\n" + "="*60)
print(f"📊 SUMMARY")
print(f"="*60)
print(f"Source file: {os.path.basename(network_inputs_file)}")
print(f"Target file: {os.path.basename(mop_temp_file)}")
print(f"Total replacements {'(preview)' if args.preview else ''}: {replacement_count}")
print(f"Sheets processed: {len(sheets_to_process)}")
print(f"IP mappings used: {len(ip_to_description)}")

if replacement_count > 0 and args.verbose:
    print(f"\n📝 Detailed change log:")
    for i, change in enumerate(changes_log, 1):
        print(f"\n{i}. Sheet: '{change['sheet']}', Cell: {change['cell']}")
        print(f"   Original: {change['original'][:60]}{'...' if len(change['original']) > 60 else ''}")
        print(f"   Updated:  {change['new'][:60]}{'...' if len(change['new']) > 60 else ''}")
        print(f"   Replaced: {', '.join([f'{ip}→{desc}' for ip, desc in change['replacements']])}")
        if change.get('route_count_replaced') and change.get('next_cell_coord'):
            print(f"   Route count replaced in next column ({change['next_cell_coord']})")

if not args.preview and not args.overwrite:
    print(f"\n💡 Tip: Review '{os.path.basename(output_file)}' and if satisfied,")
    print(f"         you can rename it to 'MOPtemp.xlsx' to replace the original file.")
elif args.preview:
    print(f"\n💡 Run without --preview to apply these changes")
    print(f"   Add --overwrite to modify the original file directly")

print(f"\n✅ Script completed successfully!")
