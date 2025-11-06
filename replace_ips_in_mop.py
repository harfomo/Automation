from openpyxl import load_workbook
import os
import re

# =========================================================
# Configuration
# =========================================================
script_dir = os.path.dirname(os.path.abspath(__file__))
network_inputs_file = os.path.join(script_dir, "network_inputs.xlsx")
mop_temp_file = os.path.join(script_dir, "MOPtemp.xlsx")
output_file = os.path.join(script_dir, "MOPtemp_updated.xlsx")

# =========================================================
# Validation
# =========================================================
if not os.path.exists(network_inputs_file):
    print(f"❌ Error: '{network_inputs_file}' not found!")
    print("   Please run the network_device_collector.py script first.")
    raise SystemExit

if not os.path.exists(mop_temp_file):
    print(f"❌ Error: '{mop_temp_file}' not found!")
    print("   Please ensure MOPtemp.xlsx exists in the script directory.")
    raise SystemExit

# =========================================================
# Load network_inputs.xlsx and build IP -> Description mapping
# =========================================================
print("📖 Loading network_inputs.xlsx...")
network_wb = load_workbook(network_inputs_file, data_only=True)

if "Neighbors" not in network_wb.sheetnames:
    print("❌ Error: 'Neighbors' sheet not found in network_inputs.xlsx")
    raise SystemExit

neighbors_sheet = network_wb["Neighbors"]

# Build mapping: Neighbor IP -> Neighbor Description
ip_to_description = {}
replacement_count = 0

print("🔍 Building IP to Description mapping from network_inputs.xlsx Neighbors sheet...")

# Try to find the header row to determine column positions
header_row = None
neighbor_desc_col = None
neighbor_ip_col = None
device_ip_col = None

for row_idx, row in enumerate(neighbors_sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
    if row:
        for col_idx, cell_value in enumerate(row):
            if cell_value:
                cell_str = str(cell_value).strip().lower()
                if 'device ip' in cell_str:
                    device_ip_col = col_idx
                    header_row = row_idx
                elif 'neighbor description' in cell_str or 'description' in cell_str:
                    neighbor_desc_col = col_idx
                    header_row = row_idx
                elif 'neighbor ip' in cell_str or (cell_str == 'neighbor ip'):
                    neighbor_ip_col = col_idx
                    header_row = row_idx
        
        if neighbor_desc_col is not None and neighbor_ip_col is not None:
            break

# Default to standard layout if headers not found
if neighbor_desc_col is None or neighbor_ip_col is None:
    print("   ⚠️  Could not auto-detect columns, using default layout:")
    print("      Column A (index 0) = Device IP")
    print("      Column B (index 1) = Neighbor Description")
    print("      Column C (index 2) = Neighbor IP")
    device_ip_col = 0
    neighbor_desc_col = 1
    neighbor_ip_col = 2
    header_row = 1
else:
    print(f"   ✅ Auto-detected columns:")
    if device_ip_col is not None:
        print(f"      Column {chr(65 + device_ip_col)} = Device IP")
    print(f"      Column {chr(65 + neighbor_desc_col)} = Neighbor Description")
    print(f"      Column {chr(65 + neighbor_ip_col)} = Neighbor IP")

# Build the mapping
for row in neighbors_sheet.iter_rows(min_row=header_row + 1, values_only=True):
    if not row or len(row) <= max(neighbor_desc_col, neighbor_ip_col):
        continue
    
    neighbor_desc = row[neighbor_desc_col] if neighbor_desc_col < len(row) else None
    neighbor_ip = row[neighbor_ip_col] if neighbor_ip_col < len(row) else None
    
    if neighbor_ip and neighbor_desc:
        neighbor_ip_str = str(neighbor_ip).strip()
        neighbor_desc_str = str(neighbor_desc).strip()
        
        # Skip empty or placeholder values
        if neighbor_ip_str and neighbor_desc_str and neighbor_ip_str.lower() not in ['none', 'n/a', '-', 'null']:
            # Store mapping (if duplicate, keep first occurrence)
            if neighbor_ip_str not in ip_to_description:
                ip_to_description[neighbor_ip_str] = neighbor_desc_str

print(f"✅ Found {len(ip_to_description)} unique Neighbor IP -> Description mappings")

if not ip_to_description:
    print("⚠️  No valid mappings found. Exiting.")
    raise SystemExit

# Show first 10 mappings as preview
print("\n📋 Sample mappings (first 10):")
for i, (ip, desc) in enumerate(list(ip_to_description.items())[:10]):
    print(f"   {ip} → {desc}")
if len(ip_to_description) > 10:
    print(f"   ... and {len(ip_to_description) - 10} more")

# Close network_inputs workbook
network_wb.close()

# =========================================================
# Helper functions for replacements
# =========================================================
def replace_ip_in_text(text, ip, description):
    """
    Replace IP with description in text, handling both IPv4 and IPv6.
    Returns (new_text, was_replaced)
    """
    original_text = text
    
    # Create case-insensitive pattern
    pattern = re.compile(re.escape(ip), re.IGNORECASE)
    
    # Check if pattern exists
    if not pattern.search(text):
        return text, False
    
    # For IPv6 (contains colons) and IPv4, use different strategies
    if ':' in ip:
        # IPv6 - direct replacement
        text = pattern.sub(description, text)
    else:
        # IPv4 - try word boundary first
        try:
            pattern_wb = re.compile(r'\b' + re.escape(ip) + r'\b', re.IGNORECASE)
            if pattern_wb.search(text):
                text = pattern_wb.sub(description, text)
            else:
                text = pattern.sub(description, text)
        except:
            text = pattern.sub(description, text)
    
    return text, (text != original_text)

def replace_route_counts_in_next_column(next_cell_value, description, route_type):
    """
    Replace route counts in the next column with description-based labels.
    For advertised-routes: <number> routes -> {description}_adv_routes routes
    For received-routes: <number> routes -> {description}_recv_routes routes
    
    Args:
        next_cell_value: Value of the next column cell
        description: The neighbor description to use
        route_type: 'advertised' or 'received'
    
    Returns (new_text, was_replaced)
    """
    if not next_cell_value or not description:
        return next_cell_value, False
    
    text = str(next_cell_value).strip()
    original_text = text
    
    # Pattern to match: optional whitespace, digits, optional whitespace, "routes"
    # Examples: "311 routes", "  500 routes  ", "123routes"
    pattern = re.compile(r'^\s*(\d+)\s*(routes?)\s*$', re.IGNORECASE)
    match = pattern.match(text)
    
    if match:
        if route_type == 'advertised':
            new_text = f"{description}_adv_routes routes"
        elif route_type == 'received':
            new_text = f"{description}_recv_routes routes"
        else:
            return text, False
        
        return new_text, True
    
    return text, False

# =========================================================
# Load MOPtemp.xlsx and search/replace
# =========================================================
print(f"\n📖 Loading {mop_temp_file}...")
mop_wb = load_workbook(mop_temp_file)

print(f"📝 Searching and replacing IPs in all sheets of MOPtemp.xlsx...")

for sheet_name in mop_wb.sheetnames:
    sheet = mop_wb[sheet_name]
    print(f"\n  🔎 Processing sheet: '{sheet_name}'")
    sheet_replacements = 0
    
    for row in sheet.iter_rows():
        # Convert row to list so we can access cells by index
        row_cells = list(row)
        
        for cell_idx, cell in enumerate(row_cells):
            if cell.value is None:
                continue
            
            cell_value = str(cell.value)
            original_value = cell_value
            replaced_ips = []
            descriptions_used = []
            
            # Step 1: Replace IPs with descriptions
            for ip, description in ip_to_description.items():
                new_value, was_replaced = replace_ip_in_text(cell_value, ip, description)
                if was_replaced:
                    cell_value = new_value
                    replaced_ips.append((ip, description))
                    descriptions_used.append(description)
            
            # Step 2: Check if this cell has advertised-routes or received-routes
            # If yes, replace route count in the NEXT column
            route_count_replaced = False
            next_cell_modified = False
            if descriptions_used:
                description = descriptions_used[0]
                route_type = None
                
                if 'advertised-routes' in cell_value.lower():
                    route_type = 'advertised'
                elif 'received-routes' in cell_value.lower():
                    route_type = 'received'
                
                # If we found a route command, check the next column
                if route_type and cell_idx + 1 < len(row_cells):
                    next_cell = row_cells[cell_idx + 1]
                    if next_cell.value is not None:
                        new_next_value, next_cell_modified = replace_route_counts_in_next_column(
                            next_cell.value, description, route_type
                        )
                        if next_cell_modified:
                            next_cell.value = new_next_value
                            route_count_replaced = True
            
            # If the value changed, update the cell
            if cell_value != original_value:
                cell.value = cell_value
                sheet_replacements += 1
                replacement_count += 1
                
                # Show details for first few replacements per sheet
                if sheet_replacements <= 5:
                    print(f"    ✓ Cell {cell.coordinate}:")
                    print(f"      Before: {original_value[:70]}{'...' if len(original_value) > 70 else ''}")
                    print(f"      After:  {cell_value[:70]}{'...' if len(cell_value) > 70 else ''}")
                    for ip, desc in replaced_ips:
                        print(f"      • Replaced IP: {ip} → {desc}")
                    if route_count_replaced:
                        next_coord = row_cells[cell_idx + 1].coordinate
                        print(f"      • Replaced route count in next column ({next_coord})")
    
    if sheet_replacements > 0:
        print(f"  ✅ Made {sheet_replacements} replacement(s) in '{sheet_name}'")
        if sheet_replacements > 5:
            print(f"     (showing first 5 replacements)")
    else:
        print(f"  ℹ️  No replacements needed in '{sheet_name}'")

# =========================================================
# Save the updated workbook
# =========================================================
print(f"\n💾 Saving updated workbook to '{output_file}'...")
mop_wb.save(output_file)

print(f"\n" + "="*60)
print(f"📊 SUMMARY")
print(f"="*60)
print(f"✅ Script completed successfully!")
print(f"📊 Total replacements made: {replacement_count}")
print(f"📁 Updated file saved as: {output_file}")
print(f"📋 Sheets processed: {len(mop_wb.sheetnames)}")
print(f"🔗 IP mappings used: {len(ip_to_description)}")
print(f"\n💡 Tip: Review '{os.path.basename(output_file)}' and if satisfied,")
print(f"         you can rename it to 'MOPtemp.xlsx' to replace the original file.")
